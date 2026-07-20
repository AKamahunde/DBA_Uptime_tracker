import openpyxl
import sys
import os

def extract_layout():
    file_path = r"e:\dba\CascadeProjects\windsurf-project\dba uptime look.xlsx"
    out_path = r"e:\dba\CascadeProjects\windsurf-project\excel_layout_report.txt"
    
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"Error opening Excel file: {e}")
        return

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("=== EXCEL LAYOUT REPORT ===\n\n")
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            f.write(f"SHEET: {sheet_name}\n")
            f.write(f"Frozen Panes: {ws.freeze_panes}\n")
            
            # Find the header row (assume first non-empty row)
            header_row = None
            headers = []
            for r in range(1, 10):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                if any(row_vals):
                    header_row = r
                    headers = row_vals
                    break
            
            if not header_row:
                f.write("  No data found.\n\n")
                continue
                
            f.write(f"  Headers (Row {header_row}):\n")
            
            for c, header_name in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=c)
                fill_color = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else "None"
                font_color = cell.font.color.rgb if cell.font and cell.font.color else "None"
                font_bold = cell.font.bold if cell.font else False
                f.write(f"    - Col {openpyxl.utils.get_column_letter(c)}: {header_name} ")
                f.write(f"(Fill: {fill_color}, Font Color: {font_color}, Bold: {font_bold})\n")
                
            f.write("\n  Sample Data (Next 2 rows):\n")
            for r in range(header_row + 1, min(header_row + 3, ws.max_row + 1)):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                f.write(f"    Row {r}: {row_vals}\n")
                
            f.write("-" * 40 + "\n\n")

    print(f"Layout report generated at: {out_path}")

if __name__ == "__main__":
    extract_layout()
