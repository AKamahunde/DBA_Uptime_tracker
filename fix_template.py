import re

with open('e:\\dba\\CascadeProjects\\windsurf-project\\create_excel_template.py', 'r', encoding='utf-8') as f:
    content = f.read()

# 1. Rename sheets
content = content.replace("'name': 'Incidents'", "'name': 'Incident Register'")
content = content.replace("'name': 'Replication'", "'name': 'Replication Monitor'")
content = content.replace("'name': 'Backups'", "'name': 'Backup & Recovery'")
content = content.replace("'name': 'Server_Health'", "'name': 'Server Health'")
content = content.replace("'name': 'Maintenance'", "'name': 'Maintenance Register'")
content = content.replace("'name': 'Capacity'", "'name': 'Capacity Planning'")
content = content.replace("'name': 'Availability'", "'name': 'Availability Calculator'")

import_dashboard = """
# ── Dashboard sheet ──────────────────────────────────────────────────────────
def create_dashboard(wb):
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_properties.tabColor = C['navy']
    ws.sheet_view.showGridLines = False

    def hdr(cell_ref, value, bg=C['navy'], fg=C['white'], bold=True, size=9, h_align='center'):
        c = ws[cell_ref]
        c.value = value
        c.font = _font(bold=bold, color=fg, size=size)
        c.fill = _fill(bg)
        c.alignment = _align(h_align, 'center', wrap=True)
        c.border = _border()
        return c

    def kpi(cell_ref, formula, num_fmt='0.00', bg=C['off_white'], fg='000000'):
        c = ws[cell_ref]
        c.value = formula
        c.font = _font(bold=False, size=10, color=fg)
        c.fill = _fill(bg)
        c.alignment = _align('center', 'center')
        c.number_format = num_fmt
        c.border = _border()
        return c

    col_widths = {'A': 18, 'B': 14, 'C': 12, 'D': 12, 'E': 12, 'F': 12,
                  'G': 12, 'H': 12, 'I': 14, 'J': 14, 'K': 14, 'L': 14,
                  'M': 12, 'N': 12, 'O': 12}
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # Row 1
    ws.merge_cells('A1:O1')
    t = ws['A1']
    t.value = 'DBA Operations Dashboard'
    t.font = _font(bold=True, size=14, color=C['white'], name='Bookman Old Style')
    t.fill = _fill(C['navy'])
    t.alignment = _align('center', 'center')
    ws.row_dimensions[1].height = 24

    # Row 2
    ws['A2'] = 'Select Database:'
    ws['A2'].font = _font(bold=True, size=11, name='Bookman Old Style')
    ws['A2'].alignment = _align('left', 'center')
    ws.row_dimensions[2].height = 20

    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"All Databases,Oracle,MySQL,MS SQL Server"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws['B2'])
    ws['B2'] = 'MySQL'
    ws['B2'].font = _font(bold=True, size=11, name='Bookman Old Style')
    ws['B2'].border = _border()
    ws['B2'].alignment = _align('left', 'center')

    # Row 3
    ws.merge_cells('A3:J3')
    ws['A3'] = 'Side-by-side comparison across all monitored databases, with detail for the selected database below.'
    ws['A3'].font = _font(italic=True, color=C['grey'], size=10, name='Bookman Old Style')
    ws['A3'].alignment = _align('left', 'center')
    ws.row_dimensions[3].height = 16

    # Row 5 Headers
    ws.row_dimensions[5].height = 40
    headers_top = [
        ('A', 'Database'), ('B', 'Overall\\nAvailability %'), ('C', 'Planned\\nDowntime\\n(hrs)'),
        ('D', 'Unplanned\\nDowntime\\n(hrs)'), ('E', 'Total\\nIncidents'), ('F', 'MTTR (mins)'),
        ('G', 'Backup\\nSuccess %'), ('H', 'Restore\\nSuccess %'), ('I', 'Replication\\nStatus'),
        ('J', 'Failed\\nBackups')
    ]
    for col, lbl in headers_top:
        hdr(f'{col}5', lbl)

    # Rows 6-8 Data
    dbs = [(6, 'Oracle'), (7, 'MySQL'), (8, 'MS SQL Server')]
    for r, db in dbs:
        ws[f'A{r}'] = db
        ws[f'A{r}'].font = _font(bold=False, name='Bookman Old Style')
        ws[f'A{r}'].border = _border()
        ws[f'A{r}'].alignment = _align('center', 'center')
        
        kpi(f'B{r}', f'=IFERROR(AVERAGEIFS(\\'Availability Calculator\\'!H:H,\\'Availability Calculator\\'!C:C,$A{r},\\'Availability Calculator\\'!A:A,TEXT(TODAY(),"YYYY-MM")),1)', '0.00%')
        kpi(f'C{r}', f'=SUMIFS(\\'Maintenance Register\\'!F:F,\\'Maintenance Register\\'!A:A,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Maintenance Register\\'!K:K,$A{r})', '0.00')
        kpi(f'D{r}', f'=SUMIFS(\\'Incident Register\\'!E:E,\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!F:F,"No",\\'Incident Register\\'!M:M,$A{r})', '0.00')
        kpi(f'E{r}', f'=COUNTIFS(\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!M:M,$A{r})', '0')
        kpi(f'F{r}', f'=IFERROR(AVERAGEIFS(\\'Incident Register\\'!E:E,\\'Incident Register\\'!M:M,$A{r},\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))*60,"No Data")', '0.00')
        kpi(f'G{r}', f'=IF(COUNTIFS(\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!B:B,$A{r},\\'Backup & Recovery\\'!C:C,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))=0,"No Data",COUNTIFS(\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!K:K,"Success",\\'Backup & Recovery\\'!B:B,$A{r},\\'Backup & Recovery\\'!C:C,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))/COUNTIFS(\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!B:B,$A{r},\\'Backup & Recovery\\'!C:C,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1)))', '0.00%')
        kpi(f'H{r}', f'=IF(COUNTIFS(\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!B:B,$A{r})=0,"No Data",COUNTIFS(\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!K:K,"Pass",\\'Backup & Recovery\\'!B:B,$A{r})/COUNTIFS(\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!B:B,$A{r}))', '0.00%')
        kpi(f'I{r}', f'=IF(COUNTIFS(\\'Replication Monitor\\'!J:J,$A{r})=0,"No Data",IF(COUNTIFS(\\'Replication Monitor\\'!J:J,$A{r},\\'Replication Monitor\\'!H:H,"Failed")>0,"Failed",IF(COUNTIFS(\\'Replication Monitor\\'!J:J,$A{r},\\'Replication Monitor\\'!H:H,"Degraded")>0,"Degraded","Healthy")))', '@')
        kpi(f'J{r}', f'=COUNTIFS(\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!K:K,"Failed",\\'Backup & Recovery\\'!B:B,$A{r},\\'Backup & Recovery\\'!C:C,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))', '0')

    ws.row_dimensions[11].height = 20
    ws['A11'] = 'Database Detail'
    ws['A11'].font = _font(bold=True, size=12, name='Bookman Old Style')

    ws.row_dimensions[12].height = 16
    ws['A12'] = 'Showing: '
    ws['A12'].font = _font(bold=True, size=10, name='Bookman Old Style')
    ws['B12'] = '=$B$2'
    ws['B12'].font = _font(bold=True, size=10, name='Bookman Old Style')

    # Row 14 Headers
    ws.row_dimensions[14].height = 40
    headers_bot = [
        ('A', 'Metric'), ('B', 'Overall\\nAvailability'), ('C', 'Planned\\nDowntime'),
        ('D', 'Unplanned\\nDowntime'), ('E', 'Total\\nIncidents'), ('F', 'MTTR'),
        ('G', 'Backup\\nSuccess Rate'), ('H', 'Restore\\nSuccess Rate'), ('I', 'Replication\\nStatus'),
        ('J', 'Last\\nSuccessful\\nFull Backup'), ('K', 'Last\\nSuccessful\\nIncremental\\nBackup'),
        ('L', 'Last DR Test'), ('M', 'Failed\\nBackups'), ('N', 'Failed\\nRestore Tests'),
        ('O', 'Replication\\nLag')
    ]
    for col, lbl in headers_bot:
        hdr(f'{col}14', lbl)

    # Row 15: Current
    hdr('A15', 'Current')
    kpi('B15', '=_xlfn.IFS($B$2="All Databases",AVERAGE(B6:B8),$B$2="Oracle",B6,$B$2="MySQL",B7,$B$2="MS SQL Server",B8)', '0.00%')
    kpi('C15', '=_xlfn.IFS($B$2="All Databases",SUM(C6:C8),$B$2="Oracle",C6,$B$2="MySQL",C7,$B$2="MS SQL Server",C8)', '0.00')
    kpi('D15', '=_xlfn.IFS($B$2="All Databases",SUM(D6:D8),$B$2="Oracle",D6,$B$2="MySQL",D7,$B$2="MS SQL Server",D8)', '0.00')
    kpi('E15', '=_xlfn.IFS($B$2="All Databases",SUM(E6:E8),$B$2="Oracle",E6,$B$2="MySQL",E7,$B$2="MS SQL Server",E8)', '0')
    kpi('F15', '=_xlfn.IFS($B$2="All Databases",AVERAGE(F6:F8),$B$2="Oracle",F6,$B$2="MySQL",F7,$B$2="MS SQL Server",F8)', '0.00')
    kpi('G15', '=_xlfn.IFS($B$2="All Databases",AVERAGE(G6:G8),$B$2="Oracle",G6,$B$2="MySQL",G7,$B$2="MS SQL Server",G8)', '0.00%')
    kpi('H15', '=_xlfn.IFS($B$2="All Databases",AVERAGE(H6:H8),$B$2="Oracle",H6,$B$2="MySQL",H7,$B$2="MS SQL Server",H8)', '0.00%')
    kpi('I15', '=IF($B$2="All Databases",IF(COUNTIF(I6:I8,"Failed")>0,"Failed",IF(COUNTIF(I6:I8,"Degraded")>0,"Degraded","Healthy")),_xlfn.IFS($B$2="Oracle",I6,$B$2="MySQL",I7,$B$2="MS SQL Server",I8))', '@')
    kpi('J15', '=IF($B$2="All Databases",_xlfn.MAXIFS(\\'Backup & Recovery\\'!C:C,\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!E:E,"Full",\\'Backup & Recovery\\'!K:K,"Success"),_xlfn.MAXIFS(\\'Backup & Recovery\\'!C:C,\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!E:E,"Full",\\'Backup & Recovery\\'!K:K,"Success",\\'Backup & Recovery\\'!B:B,$B$2))', 'M/D/YYYY')
    kpi('K15', '=IF($B$2="All Databases",_xlfn.MAXIFS(\\'Backup & Recovery\\'!C:C,\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!E:E,"Incremental",\\'Backup & Recovery\\'!K:K,"Success"),_xlfn.MAXIFS(\\'Backup & Recovery\\'!C:C,\\'Backup & Recovery\\'!A:A,"Backup",\\'Backup & Recovery\\'!E:E,"Incremental",\\'Backup & Recovery\\'!K:K,"Success",\\'Backup & Recovery\\'!B:B,$B$2))', 'M/D/YYYY')
    kpi('L15', '=IF($B$2="All Databases",_xlfn.MAXIFS(\\'Backup & Recovery\\'!C:C,\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!K:K,"Pass"),_xlfn.MAXIFS(\\'Backup & Recovery\\'!C:C,\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!K:K,"Pass",\\'Backup & Recovery\\'!B:B,$B$2))', 'M/D/YYYY')
    kpi('M15', '=_xlfn.IFS($B$2="All Databases",SUM(J6:J8),$B$2="Oracle",J6,$B$2="MySQL",J7,$B$2="MS SQL Server",J8)', '0')
    kpi('N15', '=IF($B$2="All Databases",COUNTIFS(\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!K:K,"Fail"),COUNTIFS(\\'Backup & Recovery\\'!A:A,"Restore Test",\\'Backup & Recovery\\'!K:K,"Fail",\\'Backup & Recovery\\'!B:B,$B$2))', '0')
    kpi('O15', '=IF($B$2="All Databases",MAX(\\'Replication Monitor\\'!F:F),_xlfn.MAXIFS(\\'Replication Monitor\\'!F:F,\\'Replication Monitor\\'!J:J,$B$2))', '0')

    # Row 16: Target
    hdr('A16', 'Target')
    targets = [('B','>= 99.00%'), ('C','<= 48 hrs'), ('D','<= 2 hrs'), ('E','<= 4/month'),
               ('F','<= 4 hrs'), ('G','>= 95%'), ('H','>= 80%'), ('I','Healthy'),
               ('J','< 7 days'), ('K','< 24 hrs'), ('L','< 6 months'), ('M','0'), ('N','0'), ('O','< 30 seconds')]
    for c, v in targets:
        kpi(f'{c}16', v, '@')

    # Row 17: Status
    hdr('A17', 'Status')
    kpi('B17', '=IF(B15>=0.99,"Meets SLA","Below SLA")','@')
    kpi('C17', '=IF(C15<=48,"Meets SLA","Below SLA")','@')
    kpi('D17', '=IF(D15<=2,"Meets SLA","Below SLA")','@')
    kpi('E17', '=IF(E15<=4,"Meets SLA","Below SLA")','@')
    kpi('F17', '=IF(F15<=240,"Meets SLA","Below SLA")','@')
    kpi('G17', '=IF(G15>=0.95,"Meets SLA","Below SLA")','@')
    kpi('H17', '=IF(H15>=0.80,"Meets SLA","Below SLA")','@')
    kpi('I17', '=IF(I15="Healthy","Meets SLA","Below SLA")','@')
    kpi('J17', '=IF(TODAY()-J15<=7,"Meets SLA","Below SLA")','@')
    kpi('K17', '=IF(TODAY()-K15<=1,"Meets SLA","Below SLA")','@')
    kpi('L17', '=IF(TODAY()-L15<=180,"Meets SLA","Below SLA")','@')
    kpi('M17', '=IF(M15=0,"Meets SLA","Below SLA")','@')
    kpi('N17', '=IF(N15=0,"Meets SLA","Below SLA")','@')
    kpi('O17', '=IF(O15<30,"Meets SLA","Below SLA")','@')

    # Incident Status Breakdown
    ws.row_dimensions[20].height = 20
    ws['A20'] = 'Incident Status Breakdown'
    ws['A20'].font = _font(bold=True, size=12, name='Bookman Old Style')

    ws.row_dimensions[21].height = 16
    ws['A21'] = '="This month, "&$B$2'
    ws['A21'].font = _font(italic=True, size=10, name='Bookman Old Style')

    inc_headers = ['Database', 'Open', 'In Progress', 'Resolved', 'Closed', 'On Hold', 'Total']
    for idx, header in enumerate(inc_headers):
        col = chr(ord('A') + idx)
        hdr(f'{col}23', header)
    
    kpi('A24', '=$B$2', '@')
    kpi('B24', '=COUNTIFS(\\'Incident Register\\'!L:L,"Open",\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!M:M,$B$2)', '0')
    kpi('C24', '=COUNTIFS(\\'Incident Register\\'!L:L,"In Progress",\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!M:M,$B$2)', '0.00')
    kpi('D24', '=COUNTIFS(\\'Incident Register\\'!L:L,"Resolved",\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!M:M,$B$2)', '0')
    kpi('E24', '=COUNTIFS(\\'Incident Register\\'!L:L,"Closed",\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!M:M,$B$2)', '0')
    kpi('F24', '=COUNTIFS(\\'Incident Register\\'!L:L,"On Hold",\\'Incident Register\\'!B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),\\'Incident Register\\'!M:M,$B$2)', '0.00')
    kpi('G24', '=SUM(B24:F24)', '0')

    # Footer
    ws['A27'] = 'Last Updated'
    ws['A27'].font = _font(italic=True, name='Bookman Old Style')
    ws['C27'] = '=NOW()'
    ws['C27'].number_format = 'M/D/YYYY HH:MM'
    ws['C27'].font = _font(italic=True, name='Bookman Old Style')

    # Key
    ws['A30'] = 'Key'
    ws['A30'].font = _font(bold=True, name='Bookman Old Style')
    
    ws['A31'] = 'Color'
    ws['A31'].font = _font(bold=True, name='Bookman Old Style')
    ws['A31'].border = _border()
    ws['B31'] = 'Grading Scale'
    ws['B31'].font = _font(bold=True, name='Bookman Old Style')
    ws['B31'].border = _border()

    ws['A32'].fill = _fill(C['amber_light'])
    ws['A32'].border = _border()
    ws['B32'] = 'Marginal\\nPerformance'
    ws['B32'].font = _font(name='Bookman Old Style')
    ws['B32'].border = _border()
    ws['B32'].alignment = _align('left', 'center', True)
    
    ws['A33'].fill = _fill(C['green_light'])
    ws['A33'].border = _border()
    ws['B33'] = 'Meets SLA\\nPerformance'
    ws['B33'].font = _font(name='Bookman Old Style')
    ws['B33'].border = _border()
    ws['B33'].alignment = _align('left', 'center', True)

    # Conditional Formatting for Dashboard
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    def cf(rng, val, bg, fg, op='equal'):
        ws.conditional_formatting.add(rng, CellIsRule(operator=op, formula=[f'"{val}"'], fill=_fill(bg), font=_font(bold=False, color=fg)))

    def cf_num(rng, val, bg, fg, op='greaterThanOrEqual'):
        ws.conditional_formatting.add(rng, CellIsRule(operator=op, formula=[f'{val}'], fill=_fill(bg), font=_font(bold=False, color=fg)))

    # Availability
    cf_num('B6:B8', '0.99', C['green_light'], C['green'])
    cf_num('B6:B8', '0.97', C['amber_light'], C['amber'], 'between')
    ws.conditional_formatting.add('B6:B8', CellIsRule(operator='between', formula=['0.97', '0.9899'], fill=_fill(C['amber_light']), font=_font(bold=False, color=C['amber'])))
    cf_num('B6:B8', '0.97', C['red_light'], C['red'], 'lessThan')
    cf_num('B15', '0.99', C['green_light'], C['green'])
    cf_num('B15', '0.97', C['red_light'], C['red'], 'lessThan')

    # Status text
    for rng in ['I6:I8', 'I15', 'B17:O17']:
        cf(rng, 'Healthy', C['green_light'], C['green'])
        cf(rng, 'Failed', C['red_light'], C['red'])
        cf(rng, 'Degraded', C['amber_light'], C['amber'])
        cf(rng, 'Meets SLA', C['green_light'], C['green'])
        cf(rng, 'Below SLA', C['red_light'], C['red'])
        cf(rng, 'Marginal', C['amber_light'], C['amber'])

    # Others
    cf_num('M15:N15', '0', C['green_light'], C['green'])
    cf_num('M15:N15', '1', C['red_light'], C['red'], 'greaterThanOrEqual')
    cf_num('J6:J8', '0', C['green_light'], C['green'])
    cf_num('J6:J8', '1', C['red_light'], C['red'], 'greaterThanOrEqual')

    # Incident Breakdown CF
    cf_num('B24:G24', '0', C['green_light'], C['green'])
    cf_num('B24:G24', '1', C['red_light'], C['red'], 'greaterThanOrEqual')

"""

import re
content = re.sub(r'# ── Dashboard sheet ──────────────────────────────────────────────────────────\ndef create_dashboard\(wb\):.*?# ── Build a data sheet ────────────────────────────────────────────────────────', import_dashboard + "\n\n# ── Build a data sheet ────────────────────────────────────────────────────────", content, flags=re.DOTALL)

with open('e:\\dba\\CascadeProjects\\windsurf-project\\create_excel_template.py', 'w', encoding='utf-8') as f:
    f.write(content)
